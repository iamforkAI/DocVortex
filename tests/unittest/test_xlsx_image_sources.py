"""回归 #37：openpyxl 备用图片来源必须保留原始字节和各自的 anchor。"""

from io import BytesIO
import base64
from pathlib import Path
import re
from unittest.mock import Mock
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from PIL import Image

import docvortex
from docvortex.analyzers.native.office.errors import LegacyOfficeResourceLimitError
from docvortex.analyzers.native.office.xlsx import xlsx_converter as module
from docvortex.analyzers.native.office.xlsx.ooxml_ole import workbook_sheet_parts
from docvortex.analyzers.native.office.xlsx.xlsx_converter import XlsxConverter


def _image(index: int, image_format: str = "PNG") -> bytes:
    """生成颜色和尺寸均可区分的位图载荷。"""
    stream = BytesIO()
    Image.new("RGB", (300 - index * 40, 200 - index * 20), (220 - index * 60, 30, 30 + index * 60)).save(stream, image_format)
    return stream.getvalue()


def _workbook(formats: list[str], *, placement: str = "separate") -> bytes:
    """构造多种 anchor 的 XLSX，并保留 BMP/TIFF 原始载荷供两个来源比较。"""
    workbook = Workbook()
    replacements = {}
    for index, image_format in enumerate(formats):
        sheet = workbook.active if placement != "sheets" or index == 0 else workbook.create_sheet()
        sheet["A1"] = "before images"
        sheet["A31"] = "after images"
        image = XlsxImage(BytesIO(_image(index)))
        if placement == "absolute":
            image.anchor = AbsoluteAnchor()
            image.anchor.pos.y = index * 1_000_000
            image.anchor.ext.cx = 2_000_000
            image.anchor.ext.cy = 1_000_000
        else:
            image.anchor = "C3" if placement in {"same", "sheets"} or index == 0 else f"C{20 + index - 1}"
        sheet.add_image(image)
        # 避开 openpyxl 写出时把 BMP/TIFF 转成 PNG，以测试实际包内的原始格式。
        replacements[f"xl/media/image{index + 1}.png"] = _image(0 if placement == "identical" else index, image_format)
    original = BytesIO()
    workbook.save(original)
    output = BytesIO()
    with ZipFile(original) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
    return output.getvalue()


def _collect(package: bytes, *, primary: bool = True) -> list[list]:
    """直接调用真实双来源收集路径，保留每个工作表的图片列表。"""
    converter = XlsxConverter()
    converter.workbook = load_workbook(BytesIO(package), rich_text=True)
    with ZipFile(BytesIO(package)) as archive:
        converter.zf = archive
        converter._sheet_part_by_title = workbook_sheet_parts(archive) if primary else {}
        return [converter._collect_sheet_images(sheet) for sheet in converter.workbook.worksheets]


@pytest.mark.parametrize(
    "formats",
    [["PNG"], ["PNG", "PNG"], ["PNG"] * 3, ["PNG", "JPEG"], ["JPEG", "PNG"], ["BMP", "BMP"], ["TIFF", "TIFF"]],
)
@pytest.mark.parametrize("primary", [True, False])
def test_images_match_original_content_and_anchors(formats: list[str], primary: bool) -> None:
    """主来源存在或缺失时都应逐张保留正确尺寸与原始位置。"""
    images = _collect(_workbook(formats), primary=primary)[0]
    assert len(images) == len(formats)
    for index, image in enumerate(images):
        assert image.anchor == ((2 if index == 0 else 18 + index), 2)
        payload = base64.b64decode(image.image_base64.split(",", 1)[1])
        with Image.open(BytesIO(payload)) as decoded:
            assert decoded.size == (300 - index * 40, 200 - index * 20)


@pytest.mark.parametrize("placement", ["sheets", "same", "identical", "absolute"])
def test_valid_multiple_image_placements_are_preserved(placement: str) -> None:
    """跨表、同位置不同图、同图不同位置及绝对定位均不应被误去重。"""
    images = _collect(_workbook(["PNG", "PNG"], placement=placement))
    assert sum(map(len, images)) == 2
    if placement == "sheets":
        assert list(map(len, images)) == [1, 1]
    elif placement == "same":
        assert [image.anchor for image in images[0]] == [(2, 2), (2, 2)]
    elif placement == "absolute":
        assert [image.anchor for image in images[0]] == [(None, None), (None, None)]


def test_raw_stream_is_bounded_reusable_and_has_no_fake_part_name(monkeypatch: pytest.MonkeyPatch) -> None:
    """读取 BMP 原始流时保留游标和生命周期，且绝不信任 Image.path。"""
    payload = _image(0, "BMP")
    source = BytesIO(payload)
    image = XlsxImage(source)
    source.seek(7)
    image._data = Mock(side_effect=AssertionError("原始流可读时不应转码"))
    converter = XlsxConverter()
    monkeypatch.setattr(converter, "_read_xlsx_image_member", Mock(side_effect=AssertionError("不能使用写出路径")))
    for _ in range(2):
        assert converter._raw_sheet_image(image) == (payload, None, "image/bmp")
        assert source.tell() == 7
        assert not source.closed
    monkeypatch.setattr(module, "MAX_ENTRY_BYTES", len(payload) - 1)
    with pytest.raises(LegacyOfficeResourceLimitError, match="max_entry_bytes"):
        converter._raw_sheet_image(image)
    assert source.tell() == 7
    assert not source.closed


def test_fallback_content_type_matches_transcoded_bytes() -> None:
    """缺少原始字节流时允许 _data 转码，并以输出字节识别 PNG 媒体类型。"""
    image = Mock(ref=None, format="bmp")
    payload = _image(0)
    image._data.return_value = payload
    assert XlsxConverter()._raw_sheet_image(image) == (payload, None, "image/png")
    image._data.assert_called_once()


@pytest.mark.parametrize("error", [LegacyOfficeResourceLimitError("limit"), MemoryError("memory")])
def test_original_stream_errors_do_not_trigger_unbounded_fallback(error: Exception) -> None:
    """原始流的资源错误不能被 ValueError 捕获后转入备用转码。"""
    image = Mock()
    image.ref.tell.return_value = 7
    image.ref.read.side_effect = error
    with pytest.raises(type(error)) as caught:
        XlsxConverter()._raw_sheet_image(image)
    assert caught.value is error
    image._data.assert_not_called()
    image.ref.seek.assert_called_with(7)


@pytest.mark.parametrize("error", [LegacyOfficeResourceLimitError("limit"), MemoryError("memory")])
def test_fallback_does_not_swallow_resource_errors(monkeypatch: pytest.MonkeyPatch, error: Exception) -> None:
    """备用图片来源的资源错误仍向上传播，避免用宽泛捕获绕过限额。"""
    converter = XlsxConverter()
    converter.workbook = Workbook()
    converter.workbook.active._images.append(Mock())
    monkeypatch.setattr(converter, "_raw_sheet_image", Mock(side_effect=error))
    with pytest.raises(type(error)) as caught:
        converter._collect_sheet_images(converter.workbook.active)
    assert caught.value is error


def test_fallback_decoder_receives_no_invented_source_path(monkeypatch: pytest.MonkeyPatch) -> None:
    """备用读取向公式解码和序列化传入空 part，原始图片仍能正常输出。"""
    converter = XlsxConverter()
    converter.workbook = load_workbook(BytesIO(_workbook(["PNG", "PNG"])))
    decode = Mock(return_value=None)
    serialize = Mock(return_value="data:image/png;base64,test")
    converter._image_equation_decoder = Mock(decode=decode)
    monkeypatch.setattr(module, "serialize_office_image", serialize)
    assert len(converter._collect_sheet_images(converter.workbook.active)) == 2
    assert all(call.kwargs["part_name"] is None for call in decode.call_args_list + serialize.call_args_list)


@pytest.mark.parametrize("image_format", ["PNG", "BMP"])
def test_public_markdown_export_has_exact_image_sequence(tmp_path: Path, image_format: str) -> None:
    """公开转换入口输出两张正确图片，且正文在转换后仍然保留。"""
    output = tmp_path / "images.md"
    docvortex.convert(_workbook([image_format] * 2), output, file_suffix="xlsx")
    text = output.read_text(encoding="utf-8")
    refs = re.findall(r'(?:\]\(|src=")([^)"]+)', text)
    assert len(refs) == 2
    sizes = []
    for ref in refs:
        with Image.open(output.parent / ref) as image:
            sizes.append(image.size)
    assert sizes == [(300, 200), (260, 180)]
    assert "before images" in text and "after images" in text
